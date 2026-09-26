"""Извлечение текста из PDF, XLSX, DOCX, CSV и обычных текстовых файлов.

Зачем отдельный модуль: половина полезных источников — не HTML. Отчёт в PDF,
выгрузка в XLSX, договор в DOCX. Без разбора таких файлов «исследование»
упирается в первую же ссылку на документ.

Библиотеки для разбора необязательны: если pypdf не установлен, инструмент
скажет, что именно поставить, а не упадёт с ImportError.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path

from core.logging_setup import get_logger

logger = get_logger("research.docs")


class DocumentError(Exception):
    """Документ не удалось разобрать (понятная пользователю причина)."""


#: Расширение -> (человеческое имя, пакет для установки).
FORMATS = {
    ".pdf": ("PDF", "pypdf"),
    ".xlsx": ("Excel", "openpyxl"),
    ".xlsm": ("Excel", "openpyxl"),
    ".docx": ("Word", "python-docx"),
    ".csv": ("CSV", None),
    ".tsv": ("CSV", None),
    ".txt": ("текст", None),
    ".md": ("текст", None),
    ".json": ("текст", None),
}

#: Сколько строк таблицы показываем: дальше модель всё равно не удержит контекст.
MAX_TABLE_ROWS = 200


@dataclass(slots=True)
class Document:
    """Разобранный документ."""

    title: str
    text: str
    #: Сколько всего страниц/листов — чтобы честно сказать, что показано не всё.
    parts: int = 1
    truncated: bool = False


def is_document(name: str) -> bool:
    """Похоже ли имя файла на документ, который мы умеем разбирать."""
    return Path(name).suffix.lower() in FORMATS


def extract_document(data: bytes, name: str, *, max_chars: int = 20000) -> Document:
    """Достаёт текст из байтов документа. Кидает DocumentError с объяснением."""
    suffix = Path(name).suffix.lower()
    if suffix not in FORMATS:
        raise DocumentError(
            f"Формат '{suffix or 'без расширения'}' не поддерживается. "
            f"Умею: {', '.join(sorted(FORMATS))}."
        )

    if suffix == ".pdf":
        doc = _read_pdf(data, name)
    elif suffix in (".xlsx", ".xlsm"):
        doc = _read_xlsx(data, name)
    elif suffix == ".docx":
        doc = _read_docx(data, name)
    elif suffix in (".csv", ".tsv"):
        doc = _read_csv(data, name, delimiter="\t" if suffix == ".tsv" else ",")
    else:
        doc = Document(title=name, text=_decode(data))

    if len(doc.text) > max_chars:
        doc.text = doc.text[:max_chars]
        doc.truncated = True
    return doc


def _decode(data: bytes) -> str:
    """Текст в неизвестной кодировке: utf-8, затем windows-1251 (кириллица)."""
    for encoding in ("utf-8", "utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _require(module: str, suffix: str):
    """Импорт с понятным сообщением вместо ImportError."""
    try:
        return __import__(module)
    except ImportError as exc:
        _, package = FORMATS[suffix]
        raise DocumentError(
            f"Для чтения {FORMATS[suffix][0]} нужен пакет '{package}'. "
            f"Установите: pip install {package}"
        ) from exc


# ------------------------------------------------------------------- PDF


def _read_pdf(data: bytes, name: str) -> Document:
    _require("pypdf", ".pdf")
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - битый файл, а не ошибка кода
        raise DocumentError(f"Не удалось открыть PDF: {exc}") from exc

    if reader.is_encrypted:
        # Пустой пароль открывает часть «защищённых» файлов.
        try:
            reader.decrypt("")
        except Exception as exc:  # noqa: BLE001
            raise DocumentError("PDF защищён паролем — прочитать не могу.") from exc

    chunks: list[str] = []
    for index, page in enumerate(reader.pages, 1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:  # noqa: BLE001 - одна кривая страница не должна ронять весь файл
            logger.debug("Страница %s в %s не читается", index, name, exc_info=True)
            continue
        if text:
            chunks.append(f"[стр. {index}]\n{text}")

    if not chunks:
        raise DocumentError(
            "В PDF нет текстового слоя — вероятно, это скан. "
            "Нужен OCR, простым разбором такой файл не прочитать."
        )

    title = (reader.metadata or {}).get("/Title") if reader.metadata else None
    return Document(title=str(title) if title else name, text="\n\n".join(chunks), parts=len(reader.pages))


# ----------------------------------------------------------------- Excel


def _read_xlsx(data: bytes, name: str) -> Document:
    _require("openpyxl", ".xlsx")
    import openpyxl

    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"Не удалось открыть книгу Excel: {exc}") from exc

    parts: list[str] = []
    try:
        for sheet in book.worksheets:
            rows: list[str] = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_TABLE_ROWS:
                    rows.append(f"... [показаны первые {MAX_TABLE_ROWS} строк]")
                    break
                cells = ["" if cell is None else str(cell) for cell in row]
                if any(cell.strip() for cell in cells):
                    rows.append(" | ".join(cells))
            if rows:
                parts.append(f"## Лист «{sheet.title}»\n" + "\n".join(rows))
    finally:
        book.close()

    if not parts:
        raise DocumentError("Книга Excel пуста.")
    return Document(title=name, text="\n\n".join(parts), parts=len(book.worksheets))


# ------------------------------------------------------------------ Word


def _read_docx(data: bytes, name: str) -> Document:
    _require("docx", ".docx")
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"Не удалось открыть документ Word: {exc}") from exc

    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        # Заголовки размечаем как markdown: структура помогает модели.
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append(f"\n{'#' * min(int(level), 6)} {text}")
        else:
            lines.append(text)

    for table in document.tables:
        for row in table.rows[:MAX_TABLE_ROWS]:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    if not lines:
        raise DocumentError("Документ Word не содержит текста.")
    return Document(title=name, text="\n".join(lines))


# ------------------------------------------------------------------- CSV


def _read_csv(data: bytes, name: str, *, delimiter: str = ",") -> Document:
    text = _decode(data)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows: list[str] = []
    total = 0
    for index, row in enumerate(reader):
        total = index + 1
        if index < MAX_TABLE_ROWS:
            rows.append(" | ".join(row))
    if not rows:
        raise DocumentError("Файл CSV пуст.")
    if total > MAX_TABLE_ROWS:
        rows.append(f"... [показаны {MAX_TABLE_ROWS} строк из {total}]")
    return Document(title=name, text="\n".join(rows), parts=total)
